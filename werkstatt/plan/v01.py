# 2026-10-13 13:40 · v01.py: erste Fassung des Projektplans (Version 0.1) aus den Vereinbarungen im Chat
# -----------------------------------------------------------------------------------------------
# Erzeugt den Projektplan des Luftschiff-Simulators als Excel-Datei.
# Aufbau wie der Projektplan des Gasballon-Simulators (Blätter Deckblatt, Ziele, Fahrplan,
# Anforderungen, Parameter, Modelle, Schnittstellen, Risiken, Tests, Offene_Punkte, Glossar,
# Änderungen). Spätere Versionen schreiben diese Datei mit vNN.py fort.
# Aufruf:  python v01.py <Zieldatei.xlsx>
# -----------------------------------------------------------------------------------------------
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else 'ProjektPlan_Luftschiff_0.1.xlsx'
N = 'Nutzer 13.10.2026 (Version 0.1)'   # Standard-Quelle aller Einträge dieser Fassung
B = 'Burgess, Airship Design (1927)'     # Quelle für Werte aus dem Buch

# ---------- Gestaltung: einheitliche Schrift, Kopfzeilen, Rahmen -------------------------------
FONT = 'Arial'
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_title = Font(name=FONT, size=16, bold=True, color='1F4E79')
f_sub = Font(name=FONT, size=10, italic=True, color='595959')
f_in = Font(name=FONT, size=10, color='0000FF')            # blau = Eingabewert (vom Nutzer änderbar)
fill_head = PatternFill('solid', fgColor='1F4E79')
fill_alt = PatternFill('solid', fgColor='F2F6FA')
thin = Side(style='thin', color='C9D3DD')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')

wb = openpyxl.Workbook()
wb.remove(wb.active)


def table(name, title, note, headers, widths, rows, first=4, formulas=None, lists=None):
    """Legt ein Blatt mit Titel (Zeile 1), Hinweis (Zeile 2), Kopfzeile (Zeile first-1)
    und Datenzeilen an. formulas: {Spaltenindex: Funktion(zeile) -> Formel};
    lists: {Spaltenindex: [Auswahlwerte]} für Auswahllisten (Priorität, Status ...)."""
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = f_title
    ws['A2'] = note; ws['A2'].font = f_sub
    hr = first - 1
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(hr, c, h); cell.font = f_head; cell.fill = fill_head
        cell.alignment = Alignment(wrap_text=True, vertical='center'); cell.border = box
        ws.column_dimensions[L(c)].width = w
    for k, row in enumerate(rows):
        r = first + k
        for c in range(1, len(headers) + 1):
            v = row[c - 1] if c - 1 < len(row) else None
            if formulas and c in formulas: v = formulas[c](r)
            cell = ws.cell(r, c, v)
            cell.font = f_norm; cell.alignment = wrap; cell.border = box
            if k % 2: cell.fill = fill_alt
    last = first + len(rows) - 1
    ws.auto_filter.ref = f'A{hr}:{L(len(headers))}{max(last, first)}'
    ws.freeze_panes = ws.cell(first, 1)
    if lists:
        for c, vals in lists.items():
            dv = DataValidation(type='list', formula1='"' + ','.join(vals) + '"', allow_blank=True)
            ws.add_data_validation(dv); dv.add(f'{L(c)}{first}:{L(c)}{last + 200}')
    return ws


PRIO = ['Muss', 'Soll', 'Kann']
STAT = ['Offen', 'In Arbeit', 'Teilweise', 'Umgesetzt', 'Getestet', 'Zurückgestellt']

# =============================== Deckblatt ======================================================
ws = wb.create_sheet('Deckblatt')
ws.column_dimensions['A'].width = 26; ws.column_dimensions['B'].width = 110
ws['A1'] = 'Projektplan Luftschiff-Simulator'; ws['A1'].font = f_title
ws['A2'] = ('Eine App, in der man ein Luftschiff aus einer festen Auswahl wählt oder selbst auslegt '
            '(Baukasten nach den Stichpunkten des Nutzers) und es anschließend unter echtem oder gewähltem '
            'Wetter selbst fliegt. Schwerpunkt ist eine möglichst genaue Physik der Auslegung; die '
            'Umgebung bleibt schlicht wie beim Gasballon-Simulator.')
ws['A2'].alignment = wrap; ws.merge_cells('A2:B2'); ws.row_dimensions[2].height = 45; ws['A2'].font = f_norm
info = [
    ('Version', '0.1'),
    ('Stand', '13.10.2026'),
    ('Status', 'Entwurf. Entwürfe „Baukasten“ und „Fahrt“ als Bild vom Nutzer gutgeheißen; noch kein Build. '
               'Als Nächstes: Rechenkern mit Prüfung an US Navy C-7 und USS Los Angeles (Phase 1).'),
    ('Plattform', 'Web-App für Android und iOS (wie der Gasballon-Simulator), Hochformat am Handy'),
    ('Repository', 'github.com/DerFink/Luftschiff_Simulator (öffentlich); Ordner app/ für die App, werkstatt/ für Werkzeuge'),
    ('Fachliche Grundlage', B + ', als PDF beim Nutzer; Gertler, Series 58 (Rumpfformen); '
                            'für Halbstarre und Blimps ergänzend Blakemore/Pagon, Pressure Airships (1927), noch nicht beschafft'),
    ('Verwandtes Projekt', 'Gasballon-Simulator (derfink/Bo-Twins, Ordner spiel/). Dessen Stand wird nicht verändert; '
                           'Verbesserungen sollen in beide einfließen, ohne doppelt zu entwickeln (gemeinsamer Kern geplant).'),
    ('Arbeitsweise', 'Erst prüfen und fragen, erst nach Freigabe bauen. Darstellungen vorab als Bild. '
                     'Code ausführlich kommentiert, damit erkennbar ist, wofür jeder Abschnitt zuständig ist. '
                     'Projektplan nur auf Aufforderung anlegen oder fortschreiben.'),
    ('Verantwortlich', 'Nutzer (Projektinhaber); Umsetzung Claude'),
]
for k, (a, b) in enumerate(info):
    r = 4 + k
    ws.cell(r, 1, a).font = f_bold; ws.cell(r, 2, b).font = f_norm
    ws.cell(r, 2).alignment = wrap; ws.cell(r, 1).alignment = wrap

# =============================== Ziele ==========================================================
ziele = [
    ['Ziel', 'Z-01', 'Ein Luftschiff aus einer festen Auswahl wählen oder selbst bauen, speichern und anschließend fliegen.', N],
    ['Ziel', 'Z-02', 'Die physikalischen Grundlagen der Auslegung möglichst genau bestimmen (Auftrieb, Massen, Widerstand, Leistung, Festigkeit im Überschlag). Die Umgebung bleibt schlicht.', N],
    ['Ziel', 'Z-03', 'Abschätzen lassen, wie lange und wie weit ein Luftschiff unter realen Wetterbedingungen unterwegs sein kann. Die Abschätzung macht der Nutzer selbst durch Fliegen; die App liefert Eckdaten und Auswertung.', N],
    ['Ziel', 'Z-04', 'Erkenntnis, ob ein kleines bemanntes Elektroluftschiff sinnvoll ist.', 'Nutzer 12.10.2026'],
    ['Ziel', 'Z-05', 'Spielerisch zu guten Ergebnissen führen: einfache Oberfläche, zu jeder Entscheidung ein Vorschlag, von dem man frei abweichen kann; die Folgen sind sofort an den Eckdaten sichtbar.', N],
    ['Ziel', 'Z-06', 'Größenbereich vom RC-Blimp (ab 0,25 m³) bis zum großen Starrluftschiff (rund 200.000 m³); auch autonome Luftschiffe (Drohnen).', N],
    ['Ziel', 'Z-07', 'Freie Fahrt ist der Schwerpunkt; dazu historische Szenarien, Aufgaben und Missionsplanung wie in iNav oder ArduPilot.', N],
    ['Rahmen', 'R-01', 'Web-App in einer Datei wie der Gasballon-Simulator, Karte mit Leaflet und OpenStreetMap, Wetter von Open-Meteo, 3D-Vorschau mit Three.js.', 'Vorschlag Claude, vom Nutzer angenommen'],
    ['Rahmen', 'R-02', 'Nicht kommerziell, privater Gebrauch (wie Gasballon-Simulator).', 'Annahme, zu bestätigen'],
]
table('Ziele', 'Ziele und Rahmen', 'Was die App erreichen soll und unter welchen Bedingungen.',
      ['Kategorie', 'Nr.', 'Inhalt', 'Quelle'], [12, 8, 110, 30], ziele)

# =============================== Fahrplan =======================================================
fahrplan = [
    ['1', 'Rechenkern Auslegung', 'Gertler-Rumpf (Polynom 6. Grades, L/D, m, Cp, r0, r1, spitzes Heck), Volumen, Oberfläche, Auftriebsschwerpunkt; Normatmosphäre; Traggas He/H₂ mit Reinheit und Füllgrad; Massen nach Bauteilen; Widerstand mit Größenkorrektur; Leistung nach Burgess; statische Gipfelhöhe, Ballonet; Leitwerk im Überschlag nach Burgess.', 'Prüftabelle: C-7 und USS Los Angeles nachgerechnet, Abweichungen zu Burgess-Daten benannt. Node-Tests des Kerns.', '-', 'FOR-*, MAS-*, ANT-01..03, AER-*', 'In Arbeit'],
    ['2', 'Feste Modelle', 'Die 9 festen Luftschiffe mit dem Nutzer auslegen und als Speicher-Code ablegen; Annahmen (v. a. LZ 120) festhalten.', 'Alle Modelle im Kern rechenbar; Eckdaten plausibel gegen Quellen.', 'Phase 1', 'MOD-*', 'Offen'],
    ['3', 'Baukasten-Oberfläche', 'Auswahl, Schritte nach den Stichpunkten des Nutzers, Vorschläge, Eckdaten-Leiste, 3D-Vorschau (Three.js), Ergebnis mit Massenbalken, Speicher-Code laden/speichern/teilen.', 'Ein Luftschiff lässt sich am Handy bauen, speichern und wieder laden.', 'Phase 1, Entwurf als Bild', 'BAU-*, VOR-*, SPE-*', 'Offen'],
    ['4', 'Fahrt Grundlagen', 'Karte, Ort (Wunschort, GPS, Flughafen Friedrichshafen), echtes Wetter, Flugphysik (statisch und dynamisch), Bedienung (Schub, Ruder, Höhenruder, Trimm, Ballast, Gas), Zeitraffer, Seitenansicht unter 100 m, Landung, Auswertung.', 'Freie Fahrt mit einem festen Modell vom Start bis zur Landung.', 'Phase 3', 'FAH-*, STE-*, ANZ-*', 'Offen'],
    ['5', 'Wetter und Szenarien', 'Wetter ruhig/mittel/extrem je Ort aus dem Archiv; historisches Datum mit Wetter entlang der Fahrt; Szenarien Windstill, Böig, Thermik, Front, Halle.', 'Gleiches Modell bei drei Wetterlagen fliegbar; Herkunft der Daten sichtbar.', 'Phase 4', 'MET-*', 'Offen'],
    ['6', 'Assistenten und Mission', 'Höhe halten, Kurs halten, Tempomat, Position halten; Missionsplanung mit Wegpunkten, Aktionen und Failsafe (Rückkehr bei knapper Energie).', 'Stadion-Aufgabe: 90 min Position halten mit 1 kg Kamera und Rückkehr.', 'Phase 4', 'AUT-*', 'Offen'],
    ['7', 'Aufgaben und historische Szenarien', 'LZ 129 Frankfurt–Lakehurst, LZ 120 Friedrichshafen–Berlin-Staaken, Zeppelin NT Sehenswürdigkeiten am Bodensee in 1 h; Ausdauer, Zurück zum Start, A nach B.', 'Jede Aufgabe einmal vollständig gefahren und gewertet.', 'Phase 5, 6', 'AUF-*', 'Offen'],
    ['8', 'Gemeinsamer Kern mit Gasballon', 'Atmosphäre, Wetter, Gelände, Karte als gemeinsam genutzte Teile herauslösen, ohne den Gasballon-Stand zu verändern.', 'Beide Apps nutzen dieselben Grundbausteine.', 'Phase 4', 'NFA-04', 'Offen'],
]
table('Fahrplan', 'Fahrplan', 'Phasen in der vorgesehenen Reihenfolge. Status per Auswahlliste.',
      ['Phase', 'Titel', 'Inhalt', 'Ergebnis / Abnahme', 'Voraussetzung', 'Anforderungen', 'Status'],
      [7, 26, 80, 50, 20, 22, 13], fahrplan, lists={7: STAT})

# =============================== Anforderungen ==================================================
A = []
def req(i, bereich, titel, beschr, akz, prio='Muss', status='Offen', quelle=N):
    A.append([i, bereich, titel, beschr, akz, prio, status, quelle, None])

# --- Baukasten-Ablauf
req('BAU-01', 'Baukasten', 'Ablauf nach den Stichpunkten des Nutzers',
    'Die Entscheidungen folgen der Reihenfolge der Stichpunkte: Form (L/D, Gertler), Bauweise (starr, halbstarr mit Kiel, Blimp/Prall) mit Leitwerk, Traggas (He/H₂), Volumen ⇒ L und D, Hüllengewicht, Gondel (Korb, fest am Rumpf, im Rumpf), Pax, Antriebsart, Höchstgeschwindigkeit, Ballast, Schlepptau.',
    'Alle Schritte sind in dieser Reihenfolge erreichbar; jeder Schritt lässt sich später wieder öffnen.')
req('BAU-02', 'Baukasten', 'Vorschlag zu jeder Entscheidung',
    'Zu jedem Schritt schlägt die App einen plausiblen Wert vor (z. B. Elektroantrieb: Zahl, Leistung und Lage der Motoren, Propellerdurchmesser und Akku nach der Leistungsgleichung von Burgess). Der Vorschlag ist auf dem Regler markiert und mit einem Knopf wiederherstellbar.',
    'Bei jedem Schritt erscheint ein Vorschlag; „Vorschlag“ setzt den Wert zurück.')
req('BAU-03', 'Baukasten', 'Freie Abweichung ohne Sperre',
    'Der Nutzer kann jeden Wert frei ändern, auch unsinnig (z. B. viel zu kleiner Antrieb). Die App sperrt nichts, sondern zeigt Folgen und Warnungen.',
    'Ein Antrieb mit 10 % des Vorschlags lässt sich bauen und fliegen; die Eckdaten zeigen die Folgen.')
req('BAU-04', 'Baukasten', 'Eckdaten immer sichtbar',
    'Während des gesamten Baukastens sind die Eckdaten sichtbar: Volumen, Leermasse, freie Nutzlast, Gipfelhöhe, Höchst- und Reisegeschwindigkeit, Fahrdauer, Reichweite. Die Änderung durch die letzte Entscheidung wird farbig gezeigt (grün besser, rot schlechter).',
    'Jede Änderung aktualisiert die Eckdaten sofort (unter 0,1 s).')
req('BAU-05', 'Baukasten', 'Einfache Oberfläche',
    'Die App wirkt nicht kompliziert: ein Schritt pro Bildschirm, große Vorschau, Details nur auf Wunsch.', 'Abnahme durch den Nutzer am Handy.', 'Soll')
req('BAU-06', 'Baukasten', 'Warnungen aus der Physik',
    'Beispiele: Nase beult ein (Staudruck größer als Innendruck), Ballonet zu klein für die Gipfelhöhe, Hülle überlastet, Leitwerk zu klein, H₂-Austritt am Appendix nahe den Motoren.',
    'Jede Warnung nennt den Grenzwert (z. B. „ab 33 km/h“) und eine Abhilfe.')
req('BAU-07', 'Baukasten', 'Einheiten passend zur Größe',
    'Die Anzeige schaltet je nach Größe um: g, W, Wh bei RC-Blimps; kg, kW, kWh bei kleinen; t, MW bei großen Schiffen.', 'RC-Blimp zeigt Gramm, LZ 129 Tonnen.')

# --- Form
req('FOR-01', 'Form', 'Rumpf nach Gertler (Series 58)',
    'Die Rumpfkontur folgt dem Polynom 6. Grades nach Gertler mit L/D, Lage des größten Durchmessers m, Völligkeitsgrad Cp, Nasenradius r0 und Heckradius r1.',
    'Die Kontur erfüllt alle fünf Vorgaben; Volumen nach Kontur und nach Cp stimmen überein.')
req('FOR-02', 'Form', 'Drei vorgegebene Formen',
    'Drei Gertler-Formen stehen fest zur Auswahl (z. B. schlank für Starrschiffe, Standard, völlig/kurz).', 'Drei Formen wählbar, mit Vorschaubild.')
req('FOR-03', 'Form', 'Eigenes Profil',
    'Der Nutzer kann L/D, m, Cp, r0 und r1 selbst eingeben. Unmögliche Kombinationen (Kontur mit negativem Radius) werden erkannt und gemeldet.', 'Ungültige Eingabe ergibt Hinweis statt falscher Form.')
req('FOR-04', 'Form', 'Spitzes Heck',
    'Ein spitzes Heck ist wählbar (Heckradius 0).', 'Vorschau und Rechnung zeigen ein spitz zulaufendes Heck.')
req('FOR-05', 'Form', 'Abgeleitete Größen',
    'Aus Form und Volumen folgen Länge, Durchmesser, Oberfläche, Querschnittsfläche, Lage des Auftriebsschwerpunkts und Volumen hoch 2/3.', 'Werte für ein Ellipsoid stimmen mit der geschlossenen Formel überein.')

# --- Bauweise, Masse
req('MAS-01', 'Masse', 'Masse nach Bauteilen',
    'Die Leermasse wird aus Bauteilen berechnet: Hülle (Fläche × Flächengewicht), Gerüst/Kiel, Gaszellen, Ballonets, Leitwerk, Gondel, Antrieb, Energiespeicher, Elektronik, Ausrüstung. Die Anteile nach Burgess dienen als Vorschlag für große Schiffe und zur Gegenprobe.',
    'Massenbalken zeigt alle Bauteile; Summe = Leermasse.')
req('MAS-02', 'Masse', 'Flächengewicht frei änderbar',
    'Das Flächengewicht des Hüllenstoffs ist bei jedem Modell frei änderbar. Vorschlag für kleine Folienmodelle: Tritax mit 27 g/m².', 'Änderung des Flächengewichts ändert Leermasse und Nutzlast sofort.')
req('MAS-03', 'Masse', 'Bauweisen',
    'Starr (Gerüst, Gaszellen, Außenhülle), halbstarr (Kiel, Prallhülle mit Ballonet), Blimp/Prall (Hülle mit Ballonet), Prall mit offenem Appendix ohne Ballonet (Sonderfall des Nutzers).',
    'Jede Bauweise hat eigene Massen- und Druckrechnung.')
req('MAS-04', 'Masse', 'Gerüstmasse der Starrschiffe',
    'Die Gerüstmasse skaliert mit Größe und Belastung nach Burgess (Normand-Gleichung, Biegemoment) und wird an den Starrschiffen der Auswahl abgeglichen.', 'Abweichung zur Leermasse von USS Los Angeles höchstens 10 % (Ziel).')
req('MAS-05', 'Masse', 'Nutzlast-Bauteile',
    'Pax (Masse je Person einstellbar), Besatzung oder autonom (keine Besatzung, dafür Avionik), Nutzlast wie Kamera (z. B. 1 kg), Ballast, Schlepptau.', 'Autonomes Schiff ohne Gondel und Besatzung baubar.')

# --- Gas, Statik
req('GAS-01', 'Statik', 'Traggas, Reinheit, Füllgrad',
    'Helium oder Wasserstoff mit Reinheit und Füllgrad; Bruttoauftrieb aus Normatmosphäre bzw. Wetter (Druck, Temperatur), Gastemperatur mit Über- und Untertemperatur (Sonne).', 'Auftrieb je m³ stimmt mit Tabellenwerten (Burgess) überein.')
req('GAS-02', 'Statik', 'Prallhöhe und Ballonet',
    'Prallhöhe aus Füllgrad; Ballonetgröße nach Burgess (Verhältnis abwerfbares Gewicht zu Bruttoauftrieb); höchste Höhe, aus der ohne Druckverlust abgestiegen werden kann.', 'Beispiel aus Burgess (Ballonet 20 % ⇒ 7.500 ft) wird nachgerechnet.')
req('GAS-03', 'Statik', 'Offener Appendix',
    'Bei offenem Appendix entsteht Innendruck nur aus der Gassäule (rund 11 Pa je Meter Höhe bei H₂). Steigt das Schiff, entweicht Gas; sinkt es, wird die Hülle schlaff. Formhaltung der Nase gegen den Staudruck wird geprüft.', 'Warnung „Nase beult ab … km/h ein“ mit nachvollziehbarem Wert.')
req('GAS-04', 'Statik', 'Druckhaltung der Prallhülle',
    'Bei Blimp und Halbstarr hält das Ballonet mit Gebläse den Überdruck; die nötige Druckhöhe richtet sich nach dem Staudruck bei Höchstgeschwindigkeit und dem Biegemoment (Burgess Kap. VI).', 'Warnung, wenn der Überdruck nicht reicht.', 'Soll')

# --- Aerodynamik, Antrieb
req('AER-01', 'Aerodynamik', 'Widerstand',
    'Widerstand aus Volumen hoch 2/3 und Beiwert, abhängig von Form (L/D), Anbauten (Gondel, Leitwerk, Motoren) und Reynolds-Zahl (Größenkorrektur für kleine, langsame Schiffe).', 'C-7 und Los Angeles: berechnete Höchstgeschwindigkeit bei gegebener Leistung innerhalb 10 % der Burgess-Angabe (Ziel).')
req('AER-02', 'Aerodynamik', 'Dynamischer Auftrieb',
    'Anstellwinkel des Rumpfes erzeugt Auftrieb bzw. Abtrieb; damit lässt sich Schwere oder Leichte ausgleichen.', 'Anzeige „x kg schwer ⇒ Nase +y°“ plausibel.')
req('AER-03', 'Aerodynamik', 'Leitwerk im Überschlag',
    'Leitwerksfläche und Flossenlast nach Burgess (über Volumen hoch 2/3); genaue Auslegung später (AirshipLAB zurückgestellt). Zahl der Flossen 1 bis 6, Anordnung Kreuz, X oder Y. Bei nur einer Finne kein Höhenruder: Höhe dann über Schub, Ballast und Gas.', 'Eine Finne: Höhenruder-Bedienung im Fahrbildschirm ausgeblendet.')
req('ANT-01', 'Antrieb', 'Antriebsarten',
    'Elektro (Akku), Verbrenner (Kraftstoff, Schiff wird leichter), Hybrid, Solar.', 'Jede Art liefert Masse, Leistung, Energie und Verbrauch.')
req('ANT-02', 'Antrieb', 'Leistungsgleichung nach Burgess',
    'Leistung = v³ · ρ · V^(2/3) / K mit dem Gesamtwirkungsgrad K (Burgess: Starr ≈ 63,5; Blimp ≈ 40); daraus Vorschlag für Leistung und Propeller.', 'Beispielaufgaben 1 und 4 aus Burgess Kap. II werden nachgerechnet.')
req('ANT-03', 'Antrieb', 'Lage und Zahl der Motoren',
    'Zahl, Lage (an der Gondel, am Rumpf, am Heck) und Schwenkbarkeit; Propellerdurchmesser aus Leistung und Drehzahl.', 'Vorschau zeigt Motoren an der gewählten Stelle.', 'Soll')

# --- Vorschau, Speichern, Modelle
req('VOR-01', 'Vorschau', '3D-Vorschau drehbar und zoombar',
    '3D-Ansicht des Luftschiffs, dreht sich selbst; mit einem Finger in jede Richtung drehbar, mit zwei Fingern zoombar; nach kurzer Pause dreht sie wieder selbst.', 'Am Handy flüssig (mindestens 30 Bilder/s).')
req('VOR-02', 'Vorschau', 'Darstellung nach Bauweise',
    'Nur Starrluftschiffe zeigen die Längsstriche des Gerüsts auf der Hülle; Blimp und Halbstarr haben eine glatte Hülle. Gondel, Leitwerk (Zahl und Anordnung der Flossen), Motoren, Propeller und Appendix werden wie gewählt dargestellt.', 'Jede Bauweise und jede Flossenzahl sieht im Bild richtig aus.')
req('SPE-01', 'Speichern', 'Schlanker Speicher-Code',
    'Ein Luftschiff wird als kurzer, lesbarer Text gespeichert (z. B. „LS1|B:P|G:H2|V:2000|…“). Nur Entscheidungen werden gespeichert, Ergebnisse rechnet die App neu. Nicht angegebene Werte gelten als Vorschlag. Mit Versionskennung.', 'Code passt in eine Nachricht oder einen QR-Code; Laden ergibt dasselbe Schiff.')
req('SPE-02', 'Speichern', 'Eigene Luftschiffe speichern',
    'Eigene Luftschiffe lassen sich auf dem Gerät speichern, kopieren und teilen.', 'Gespeichertes Schiff nach Neustart wieder da.')
req('MOD-01', 'Modelle', 'Feste Auswahl',
    'Starr: USS Los Angeles (LZ 126), LZ 120 „Bodensee“ (kurze Fassung 1919), LZ 129 „Hindenburg“ (oder LZ 127). Halbstarr: N-1 „Norge“, Zeppelin NT. Blimp: US Navy C-7, Goodyear GZ-20 (oder Skyship 600). Klein: Kleines H₂-Luftschiff (1.000–3.000 m³, offen, länglicher Korb, 2 E-Antriebe), RC-Blimp (0,5 m³, He, Tritax).', 'Alle 9 Modelle im Blatt Modelle beschrieben und in der App wählbar.')
req('MOD-02', 'Modelle', 'Annahmen dokumentiert',
    'Wo Daten fehlen (v. a. LZ 120), werden Annahmen gemeinsam festgelegt und im Plan dokumentiert.', 'Blatt Modelle nennt Quelle oder Annahme je Wert.')

# --- Fahrt
req('FAH-01', 'Fahrt', 'Selbst fliegen',
    'Steuerung: Schub, links/rechts (Seitenruder), hoch/runter (Höhenruder), Trimm, Ballast abwerfen, Gas ablassen (wenn vorgesehen).', 'Alle Bedienelemente wirken auf die Physik.')
req('FAH-02', 'Fahrt', 'Startarten',
    'Start mit Schub, mit Ballast oder mit beidem. Vor dem Start zeigt die App die statische Bilanz (schwer/leicht).', 'Alle drei Startarten möglich.')
req('FAH-03', 'Fahrt', 'Flugphysik',
    'Statischer Auftrieb, dynamischer Auftrieb, Widerstand, Schub, Wind, Trägheit mit zusätzlicher Masse (Burgess), Gaserwärmung durch Sonne, Kraftstoffverbrauch macht leichter, Akku nicht.', 'Physik-Tests des Kerns.')
req('FAH-04', 'Fahrt', 'Zeitraffer',
    'Zeitraffer ×1, ×5, ×10, ×50, ×100, ×500, ×1000; intern kleinere Rechenschritte, damit die Physik stabil bleibt.', 'Bei ×1000 kein Aufschwingen der Physik.')
req('FAH-05', 'Fahrt', 'Seitenansicht bei der Landung',
    'Unter 100 m über Grund erscheint automatisch ein Seitenprofil des Luftschiffs mit Boden, Höhe und Sinkrate. Jederzeit ein- und ausblendbar; die Automatik lässt sich abschalten.', 'Erscheint unter 100 m; Schalter wirken.')
req('FAH-06', 'Fahrt', 'Landung',
    'Landung mit Sinkgeschwindigkeit; später Haltetrupp bzw. Mast.', 'Landegeschwindigkeit in der Auswertung.')
req('ANZ-01', 'Anzeige', 'Anzeigen während der Fahrt',
    'Höhe, Steigrate, Fahrt durch die Luft, Fahrt über Grund, Bilanz schwer/leicht mit nötigem Anstellwinkel, Füllgrad mit Prallhöhe, Energie mit Restzeit, Ballast; auf der Karte Kurs (Nase) und Weg über Grund, maßstäbliches Schiff.', 'Alle Werte sichtbar am Handy.')
req('ANZ-02', 'Anzeige', 'Rückkehr-Reichweite',
    'Auf der Karte ein Bereich, aus dem man mit der Restenergie bei aktuellem Wind noch zum Start zurückkommt, und die Uhrzeit „Umkehren bis“.', 'Bereich ändert sich mit Wind und Energie.')
req('ANZ-03', 'Anzeige', 'Auswertung',
    'Nach der Landung: Dauer, Strecke über Grund und durch die Luft, Restenergie, Restballast, Landegeschwindigkeit, Verlauf (Höhe, Energie, Geschwindigkeit), Vergleich mit den Eckdaten des Baukastens mit Begründung, gezählte Warnungen; „Zum Baukasten“.', 'Auswertung nach jeder Fahrt.')

# --- Automatik
req('AUT-01', 'Automatik', 'Assistenten',
    'Höhe halten, Kurs halten, Tempomat, Position halten.', 'Jeder Assistent hält seinen Wert bei echtem Wind.')
req('AUT-02', 'Automatik', 'Position halten (Stadion-Aufgabe)',
    'Beispiel: 90 min über einem Stadion bleiben mit 1 kg Kamera, danach zurückfliegen. Zeigt, ob ein elektrisches Luftschiff das schafft.', 'Aufgabe fahrbar und gewertet.')
req('AUT-03', 'Automatik', 'Missionsplanung',
    'Wegpunkte auf der Karte mit Höhe, Geschwindigkeit und Aktion (x min Position halten, kreisen, landen); Failsafe: Rückkehr bei Energie unter x % oder wenn die Restenergie bei aktuellem Gegenwind gerade noch reicht. Mission vorab durchrechnen und live verfolgen. Vorbild iNav/ArduPilot.', 'Mission mit drei Wegpunkten und Failsafe fliegt selbstständig.', 'Soll')

# --- Aufgaben
req('AUF-01', 'Aufgaben', 'Freie Fahrt als Schwerpunkt',
    'Freie Fahrt steht an erster Stelle.', 'Erster Eintrag im Startbildschirm.')
req('AUF-02', 'Aufgaben', 'Allgemeine Aufgaben',
    'Ausdauer, Zurück zum Start, Von A nach B, Position halten.', 'Jede Aufgabe gewertet.', 'Soll')
req('AUF-03', 'Aufgaben', 'Historische Szenarien',
    'LZ 129 Frankfurt → Lakehurst; LZ 120 Friedrichshafen → Berlin-Staaken; Zeppelin NT: möglichst viele Sehenswürdigkeiten am Bodensee in 1 h (z. B. Meersburg, Mainau, Lindau, Pfahlbauten; zählt bei Vorbeiflug nah genug und tief genug).', 'Alle drei Szenarien fahrbar.', 'Soll')

# --- Wetter, Ort
req('MET-01', 'Wetter', 'Ort wählen',
    'Wunschort auf der Karte, aktueller Ort per GPS, Standard Flughafen Friedrichshafen.', 'Alle drei Wege funktionieren.')
req('MET-02', 'Wetter', 'Echtes Wetter',
    'Aktuelles Wetter von Open-Meteo für den Ort des Schiffs, entlang der Fahrt nachgeladen.', 'Wetter wechselt bei langer Fahrt.')
req('MET-03', 'Wetter', 'Ruhig, mittel, extrem je Ort',
    'Nach Wahl des Ortes: ruhiges, mittleres oder extremes Wetter, abgeleitet aus dem Wetterarchiv des Ortes (z. B. ruhigste 10 %, typischer Tag, windigste 5 % der Tage).', 'Küstenort ergibt rauere Werte als Friedrichshafen.')
req('MET-04', 'Wetter', 'Historisches Datum',
    'Datum in der Vergangenheit wählen; wo immer das Schiff hinfliegt, wird das Wetter zu diesem Zeitpunkt an diesem Ort geladen. Archiv ab 1940; für frühere Szenarien (1919, 1936) dasselbe Datum aus einem neueren Jahr (z. B. 2025).', 'Wetter folgt Ort und Simulationszeit.')
req('MET-05', 'Wetter', 'Herkunft der Wetterdaten sichtbar',
    'Ein Feld nennt immer, woher und von wann die Daten stammen (Vorhersage, Archivdatum, Ersatzdatum, Szenario mit Ausgangstag).', 'Beispiel: „Wetter vom 04.05.2025 (Ersatz für 04.05.1936)“.')
req('MET-06', 'Wetter', 'Szenarien',
    'Windstill, Böig, Thermik, Front, Halle (ohne Wind, für RC-Blimps).', 'Jedes Szenario wählbar.', 'Soll')
req('MET-07', 'Wetter', 'Windprofil über 100 m im Archiv',
    'Das Archiv liefert Wind nur in 10 m und 100 m; darüber wird das Windprofil hochgerechnet. Hinweis in der App, dass es oberhalb von etwa 300–500 m ungenauer wird.', 'Hinweis sichtbar bei Archivwetter.')

# --- Nichtfunktional, Daten
req('NFA-01', 'Nichtfunktional', 'Web-App am Handy',
    'Läuft im Browser von Android und iOS, installierbar, Start ohne Netz (Service Worker) wie beim Gasballon.', 'Installation auf beiden Systemen.')
req('NFA-02', 'Nichtfunktional', 'Kommentierter Code',
    'Code ausführlich kommentiert, damit der Nutzer erkennt, wofür jeder Abschnitt zuständig ist.', 'Durchsicht.', quelle='Nutzer (allgemeine Vorgabe)')
req('NFA-03', 'Nichtfunktional', 'Rechenkern getrennt testbar',
    'Rechenkern zwischen Markierungen wie beim Gasballon, als core.js für Node-Tests exportierbar.', 'Node-Tests laufen.')
req('NFA-04', 'Nichtfunktional', 'Gemeinsame Teile mit Gasballon',
    'Atmosphäre, Wetter, Gelände, Karte sollen gemeinsam genutzt werden, ohne den Gasballon-Stand zu verändern.', 'Keine Änderung an Bo-Twins ohne Freigabe.', 'Soll', quelle='Nutzer 12.10.2026')
req('NFA-05', 'Nichtfunktional', 'Größenbereich',
    'Rechnung stimmt von 0,25 m³ bis rund 200.000 m³.', 'RC-Blimp und LZ 129 liefern plausible Werte.')

def fixA(r): return f'=COUNTIF(Tests!$C$4:$C$300,"*"&A{r}&"*")'
ws = table('Anforderungen', 'Anforderungen',
           'Stand 13.10.2026, noch kein Build. Priorität und Status per Auswahlliste. Spalte I zählt automatisch die Testfälle (Blatt Tests), die die Anforderung nennen.',
           ['ID', 'Bereich', 'Titel', 'Beschreibung', 'Akzeptanz', 'Prio', 'Status', 'Quelle', 'Testfälle'],
           [9, 13, 28, 70, 40, 8, 11, 22, 9], A, formulas={9: fixA}, lists={6: PRIO, 7: STAT})

# =============================== Parameter ======================================================
P = [
    ('Volumenbereich, kleinstes Schiff', 0.25, 'm³', N + ' (RC-Blimp)'),
    ('Volumenbereich, größtes Schiff', 200000, 'm³', 'LZ 129, rund'),
    ('Flächengewicht Folie Tritax (Vorschlag kleine Modelle)', 27, 'g/m²', N),
    ('Flächengewicht gummierter Baumwollstoff 1927', 475, 'g/m²', B + ': 14 oz/yd², Reißfestigkeit rund 80 lb/in'),
    ('Grenze runde Prallhülle mit Stoff von 1927', 5700, 'm³', B + ': rund 200.000 ft³'),
    ('Untere Grenze sinnvoller Starrluftschiffe', 28300, 'm³', B + ': rund 1 Mio. ft³'),
    ('Gesamtwirkungsgrad K, Starr', 63.5, '-', B + ', Kap. II Aufgabe 1'),
    ('Gesamtwirkungsgrad K, Blimp', 40, '-', B + ', Kap. II Aufgabe 4'),
    ('Feste Massen ohne Antrieb, Starr', 0.30, '× D', B + ', Kap. II (D = Masse der verdrängten Luft)'),
    ('Feste Massen ohne Antrieb, Blimp', 0.40, '× D', B + ', Kap. II'),
    ('Besatzung, Vorräte, Ballast, Starr', 0.055, '× D', B + ', Kap. II'),
    ('Besatzung, Vorräte, Ballast, Blimp', 0.10, '× D', B + ', Kap. II'),
    ('Triebwerksmasse 1927, Starr / Blimp', '8 / 6', 'lb/PS', B + ', Kap. II'),
    ('Kraftstoff 1927', 0.6, 'lb/(PS·h)', B + ', Kap. II'),
    ('Luftdichte Meereshöhe (Normatmosphäre)', 1.225, 'kg/m³', 'ISA'),
    ('Druckzunahme in H₂-Gassäule über Appendix', 11.2, 'Pa/m', 'Rechnung Claude: g·(ρLuft − ρH₂)'),
    ('Nase beult ein, kleines H₂-Schiff (Ø ≈ 10 m)', 33, 'km/h', 'Überschlag Claude, zu prüfen'),
    ('Seitenansicht automatisch unter', 100, 'm über Grund', N),
    ('Zeitraffer-Stufen', '1, 5, 10, 50, 100, 500, 1000', '×', N),
    ('Standardort', 'Flughafen Friedrichshafen (EDNY), 47,671° N 9,511° O', '-', N),
    ('Wetterarchiv Open-Meteo verfügbar ab', 1940, 'Jahr', 'open-meteo.com, Historical Weather API'),
    ('Ersatzjahr für Szenarien vor 1940', 2025, 'Jahr', N),
    ('Wind im Archiv nur in', '10 und 100', 'm', 'open-meteo.com, Historical Weather API'),
    ('Flossenzahl Leitwerk', '1 bis 6', 'Stück', N),
    ('Masse Kamera (Stadion-Aufgabe)', 1, 'kg', N),
    ('Dauer Position halten (Stadion-Aufgabe)', 90, 'min', N),
]
ws = table('Parameter', 'Parameter', 'Vorgaben und Kennwerte mit Quelle. Blau = vom Nutzer vorgegeben oder änderbar.',
           ['Parameter', 'Wert', 'Einheit', 'Quelle'], [52, 36, 14, 60], [list(p) for p in P])
for r in range(4, 4 + len(P)):
    if 'Nutzer' in str(ws.cell(r, 4).value): ws.cell(r, 2).font = f_in

# =============================== Modelle ========================================================
M = [
    ['Starr', 'USS Los Angeles (LZ 126)', 'He', '≈ 70.000', '1924', 'Burgess Abb. 1, umfangreich', 'Prüfschiff Phase 1'],
    ['Starr', 'LZ 120 „Bodensee“ (kurz)', 'H₂', '≈ 20.000', '1919', 'Wenige Daten; Annahmen gemeinsam festlegen', 'Lieblingszeppelin des Nutzers; Szenario Friedrichshafen → Staaken'],
    ['Starr', 'LZ 129 „Hindenburg“ (oder LZ 127)', 'H₂', '≈ 200.000', '1936', 'Zu beschaffen', 'Szenario Frankfurt → Lakehurst'],
    ['Halbstarr', 'N-1 „Norge“', 'H₂', '≈ 19.000', '1926', 'Burgess Abb. 2; Blakemore/Pagon', ''],
    ['Halbstarr', 'Zeppelin NT', 'He', '≈ 8.400', '1997', 'Herstellerdaten zu beschaffen', 'Szenario Sehenswürdigkeiten Bodensee'],
    ['Blimp', 'US Navy C-7', 'He', '≈ 5.200', '1921', 'Burgess Abb. 3, Druckmessungen', 'Prüfschiff Phase 1'],
    ['Blimp', 'Goodyear GZ-20 (oder Skyship 600)', 'He', '≈ 5.700', '1969', 'Zu beschaffen', ''],
    ['Klein', 'Kleines H₂-Luftschiff', 'H₂', '1.000–3.000', 'neu', 'Vorgabe Nutzer: unten offen mit kurzem Appendix, länglicher Korb, 2 E-Antriebe an der Gondel', 'Name noch offen'],
    ['Klein', 'RC-Blimp', 'He', '≈ 0,5', 'neu', 'Tritax 27 g/m²', ''],
]
table('Modelle', 'Feste Auswahl', 'Die festen Luftschiffe der App. Volumen gerundet, vor Phase 2 zu prüfen.',
      ['Bauart', 'Name', 'Gas', 'Volumen m³', 'Jahr', 'Datenlage', 'Bemerkung'], [11, 32, 6, 13, 8, 50, 45], M)

# =============================== Schnittstellen =================================================
S = [
    ['IF-01', 'App', 'Open-Meteo Vorhersage', 'HTTPS, JSON', 'Wetter für den Ort des Schiffs mit Druckflächen', 'Server zu App', 'Offen', 'CC BY 4.0; wie beim Gasballon'],
    ['IF-02', 'App', 'Open-Meteo Archiv', 'HTTPS, JSON', 'Historisches Wetter ab 1940 (Wind 10 m und 100 m, Temperatur, Druck)', 'Server zu App', 'Offen', 'Für ruhig/mittel/extrem und historisches Datum'],
    ['IF-03', 'App', 'Open-Meteo Höhenmodell', 'HTTPS, JSON', 'Geländehöhen', 'Server zu App', 'Offen', 'Wie beim Gasballon'],
    ['IF-04', 'App', 'OpenStreetMap (Leaflet)', 'HTTPS, Kacheln', 'Kartenbild', 'Server zu App', 'Offen', 'Namensnennung nötig'],
    ['IF-05', 'App', 'Three.js', 'Bibliothek', '3D-Vorschau', 'eingebunden', 'Offen', 'Fest eingebunden für Start ohne Netz'],
    ['IF-06', 'Nutzer', 'App', 'Text / QR', 'Speicher-Code eines Luftschiffs', 'beide', 'Offen', 'SPE-01'],
]
table('Schnittstellen', 'Schnittstellen', 'Externe Dienste und Datenwege.',
      ['ID', 'Von', 'Nach', 'Medium', 'Inhalt', 'Richtung', 'Status', 'Anmerkung'], [8, 12, 24, 14, 50, 14, 10, 40], S, lists={7: STAT})

# =============================== Risiken ========================================================
R = [
    ['R-01', 'Fachlich', 'Massenmodell weicht stark ab', 'Gerüst- und Hüllenmassen aus 1927 passen nicht zu modernen Werkstoffen', 'Falsche Eckdaten, falsche Schlüsse zum E-Luftschiff', 3, 4, None, None, 'Bauteilrechnung mit änderbaren Kennwerten; Abgleich an echten Schiffen', 'Claude', 'Offen'],
    ['R-02', 'Fachlich', 'Widerstand kleiner Schiffe falsch', 'Burgess-Beiwerte gelten für große Schiffe', 'RC-Blimp zu schnell', 3, 3, None, None, 'Reynolds-Korrektur; Abgleich mit Messwerten von RC-Blimps', 'Claude', 'Offen'],
    ['R-03', 'Daten', 'Wenige Daten zu LZ 120 und Norge', 'Historische Quellen lückenhaft', 'Modell ungenau', 4, 2, None, None, 'Annahmen gemeinsam festlegen und dokumentieren (MOD-02)', 'Nutzer, Claude', 'Offen'],
    ['R-04', 'Daten', 'Archivwetter ohne Höhenwind', 'Archiv nur 10 m und 100 m', 'Ungenaue Winde in größerer Höhe', 5, 2, None, None, 'Windprofil hochrechnen, Hinweis anzeigen (MET-07)', 'Claude', 'Offen'],
    ['R-05', 'Technik', 'Rechenaufwand am Handy', '3D-Vorschau, Physik und Zeitraffer ×1000', 'Ruckeln, Akkuverbrauch', 3, 3, None, None, 'Einfache Geometrie, Physik mit festen Schritten, Messung am Handy', 'Claude', 'Offen'],
    ['R-06', 'Umfang', 'Projekt wird zu groß', 'Viele Wünsche (Mission, Szenarien, Wetterarchiv)', 'Lange bis zur ersten nutzbaren Fassung', 4, 3, None, None, 'Phasen strikt nacheinander; Freie Fahrt zuerst', 'Nutzer, Claude', 'Offen'],
    ['R-07', 'Sicherheit', 'Falscher Eindruck der Machbarkeit', 'Simulation wird als Bauentscheid genommen', 'Fehlentscheidung beim echten Bau', 2, 5, None, None, 'Hinweis in der App: Abschätzung, kein Nachweis', 'Claude', 'Offen'],
]
def fH(r): return f'=F{r}*G{r}'
def fI(r): return f'=IF(H{r}>=15,"hoch",IF(H{r}>=8,"mittel","niedrig"))'
table('Risiken', 'Risiken', 'Wahrscheinlichkeit (W) und Schaden (S) je 1 bis 5; Wert = W × S (Formel); Stufe ab 15 hoch, ab 8 mittel.',
      ['ID', 'Kategorie', 'Risiko', 'Ursache', 'Auswirkung', 'W', 'S', 'Wert', 'Stufe', 'Maßnahme', 'Verantwortlich', 'Status'],
      [7, 11, 30, 34, 30, 5, 5, 7, 9, 40, 14, 10], R, formulas={8: fH, 9: fI}, lists={12: STAT})

# =============================== Tests ==========================================================
T = [
    ['T-01', 'Kern', 'FOR-01, FOR-04, FOR-05', 'Gertler-Kontur', 'Node', 'Kontur für Ellipsoid (m 0,5, r0 0,5, r1 0,5, Cp 2/3) und für ein spitzes Heck rechnen', 'Ellipsoid: Volumen und Oberfläche wie geschlossene Formel; spitzes Heck: Radius → 0', 'Offen'],
    ['T-02', 'Kern', 'GAS-01, GAS-02', 'Auftrieb und Ballonet', 'Node', 'Auftrieb je m³ He/H₂; Burgess-Beispiel Ballonet 20 %', 'Werte wie Burgess (7.500 ft)', 'Offen'],
    ['T-03', 'Kern', 'ANT-02, AER-01, MAS-01', 'Prüfung C-7', 'Node', 'C-7 mit Burgess-Daten rechnen', 'Höchstgeschwindigkeit und Leermasse innerhalb 10 %', 'Offen'],
    ['T-04', 'Kern', 'ANT-02, AER-01, MAS-04', 'Prüfung USS Los Angeles', 'Node', 'Los Angeles mit Burgess-Daten rechnen', 'Höchstgeschwindigkeit und Leermasse innerhalb 10 %', 'Offen'],
    ['T-05', 'Kern', 'ANT-02', 'Burgess Kap. II Aufgaben 1 und 4', 'Node', 'Größenbestimmung nachrechnen', 'Volumen und Leistung wie im Buch', 'Offen'],
    ['T-06', 'Kern', 'GAS-03, BAU-06', 'Offener Appendix', 'Node', 'Kleines H₂-Schiff 2.000 m³', 'Grenzgeschwindigkeit Nase ≈ 30–35 km/h', 'Offen'],
    ['T-07', 'Software', 'SPE-01, SPE-02', 'Speicher-Code', 'Handy', 'Schiff bauen, Code kopieren, in neuem Browser laden', 'Gleiches Schiff, gleiche Eckdaten', 'Offen'],
    ['T-08', 'Software', 'VOR-01, VOR-02, AER-03', '3D-Vorschau', 'Handy', 'Alle Bauweisen und Flossenzahlen 1 bis 6 ansehen, drehen, zoomen', 'Striche nur bei Starr; Flossen richtig', 'Offen'],
    ['T-09', 'Software', 'BAU-01, BAU-02, BAU-03, BAU-04, BAU-05, BAU-06, BAU-07', 'Baukasten-Ablauf', 'Handy', 'Eigenes Schiff von Grund auf bauen; Antrieb viel zu klein wählen', 'Eckdaten reagieren, Warnungen erscheinen, nichts gesperrt', 'Offen'],
    ['T-10', 'Software', 'FAH-01, FAH-02, FAH-03, FAH-04, FAH-05, FAH-06, ANZ-01, ANZ-03', 'Freie Fahrt', 'Handy', 'Kleines H₂-Schiff in Friedrichshafen starten, fahren, unter 100 m landen', 'Seitenansicht ab 100 m; Auswertung erscheint', 'Offen'],
    ['T-11', 'Software', 'AUT-01, AUT-02', 'Stadion-Aufgabe', 'Handy', '90 min Position halten mit 1 kg Kamera, dann zurück', 'Ergebnis „geschafft“ oder „Energie reicht nicht“ nachvollziehbar', 'Offen'],
    ['T-12', 'Software', 'MET-03, MET-04, MET-05, MET-07', 'Wetterarten', 'Handy', 'Ruhig/mittel/extrem und historisches Datum 1936 wählen', 'Herkunft sichtbar; Ersatzjahr genannt', 'Offen'],
]
table('Tests', 'Tests', 'Testfälle; Spalte C nennt die geprüften Anforderungen.',
      ['ID', 'Art', 'Anforderungen', 'Titel', 'Umgebung', 'Ablauf', 'Erwartung', 'Status'],
      [7, 10, 26, 26, 10, 50, 46, 10], T, lists={8: STAT})

# =============================== Offene Punkte ==================================================
O = [
    ['OP-01', 'Modelle', 'Welche Annahmen gelten für LZ 120 „Bodensee“ (kurz)? Gerüst, Gaszellen, Motoren, Leermasse.', 'Gemeinsam aus Quellen und Burgess ableiten', 'Nutzer, Claude', '2', 'Muss', '', '', 'Offen'],
    ['OP-02', 'Modelle', 'LZ 129 oder LZ 127 als drittes Starrschiff? Goodyear GZ-20 oder Skyship 600?', 'Je eines / beide', 'Nutzer', '2', 'Soll', '', '', 'Offen'],
    ['OP-03', 'Modelle', 'Name für das kleine H₂-Luftschiff?', '', 'Nutzer', '2', 'Kann', '', '', 'Offen'],
    ['OP-04', 'Form', 'Welche drei Gertler-Formen werden fest vorgegeben (m, Cp, r0, r1)?', 'Vorschlag Claude nach Phase 1', 'Nutzer', '1', 'Muss', '', '', 'Offen'],
    ['OP-05', 'Quellen', '„Pressure Airships“ (Blakemore/Pagon, 1927) für Halbstarre und Blimps beschaffen?', 'Frei verfügbare Kopie suchen / ohne', 'Nutzer', '2', 'Soll', '', '', 'Offen'],
    ['OP-06', 'Leitwerk', 'Genaue Leitwerksauslegung mit AirshipLAB (Alex)?', 'Screenshots als Gegenprobe / Code von Alex / nur Burgess', 'Nutzer', 'später', 'Kann', 'Vorerst nur Überschlag nach Burgess', '13.10.2026', 'Zurückgestellt'],
    ['OP-07', 'Technik', 'Wie werden gemeinsame Teile mit dem Gasballon geteilt (Kopie oder gemeinsames Modul)?', 'Kopie mit Abgleich / gemeinsames Modul', 'Nutzer, Claude', '8', 'Soll', '', '', 'Offen'],
    ['OP-08', 'Recht', 'Gilt auch hier: nicht kommerziell, privater Gebrauch, Hinweis in der App?', 'Ja / Nein', 'Nutzer', '3', 'Soll', '', '', 'Offen'],
    ['OP-09', 'Betrieb', 'Wo läuft die App: GitHub Pages des neuen Repositorys, später Server cloud.aerarium.de?', '', 'Nutzer', '3', 'Soll', '', '', 'Offen'],
]
table('Offene_Punkte', 'Offene Punkte', 'Fragen, die vor oder während der Umsetzung zu klären sind.',
      ['ID', 'Thema', 'Frage', 'Optionen', 'Zuständig', 'Phase', 'Prio', 'Entscheidung', 'Entschieden am', 'Status'],
      [7, 10, 60, 36, 14, 7, 7, 34, 12, 12], O, lists={7: PRIO, 10: ['Offen', 'In Klärung', 'Entschieden', 'Zurückgestellt']})

# =============================== Glossar ========================================================
G = [
    ['Starrluftschiff', 'Luftschiff mit tragendem Gerüst aus Ringen und Längsträgern; Gas in Zellen, Form unabhängig vom Gasdruck.'],
    ['Halbstarr', 'Prallhülle mit tragendem Kiel (z. B. Norge, Zeppelin NT mit Innengerüst).'],
    ['Blimp / Prallluftschiff', 'Form wird allein durch Überdruck gehalten; Ballonets gleichen Volumenänderungen aus.'],
    ['Ballonet', 'Luftsack in der Hülle, mit dem der Überdruck bei Höhenänderung gehalten wird.'],
    ['Appendix', 'Offener Füllschlauch unten an der Hülle; über ihn entweicht Gas beim Steigen (wie beim Gasballon).'],
    ['Prallhöhe', 'Höhe, in der das Gas die Hülle bzw. Zellen ganz ausfüllt; darüber geht Gas verloren.'],
    ['Gertler-Form (Series 58)', 'Familie von Rumpfformen, beschrieben durch ein Polynom 6. Grades mit L/D, m, Cp, r0, r1.'],
    ['L/D', 'Schlankheit: Länge durch größten Durchmesser.'],
    ['m', 'Lage des größten Durchmessers, bezogen auf die Länge (von der Nase).'],
    ['Cp', 'Völligkeitsgrad: Volumen durch Volumen des umschreibenden Zylinders.'],
    ['r0, r1', 'Nasen- und Heckradius, dimensionslos (Radius × L / D²).'],
    ['K (Burgess)', 'Gesamtwirkungsgrad aus Rumpfwiderstand und Propeller in der Leistungsgleichung.'],
    ['D (Burgess)', 'Masse der vom Luftschiff verdrängten Luft in der Normatmosphäre.'],
    ['Schwer / leicht', 'Unterschied zwischen Gewicht und statischem Auftrieb; wird mit dynamischem Auftrieb, Ballast oder Gas ausgeglichen.'],
    ['Dynamischer Auftrieb', 'Auftrieb des angeströmten Rumpfes bei Anstellwinkel.'],
    ['LTW', 'Leitwerk (Flossen, Seiten- und Höhenruder).'],
    ['Tritax', 'Leichte Folie für Hüllen kleiner Modelle, rund 27 g/m².'],
    ['Speicher-Code', 'Kurzer Text, der alle Entscheidungen eines Luftschiffs enthält (z. B. „LS1|B:P|G:H2|V:2000|…“).'],
    ['Failsafe', 'Selbsttätige Rückkehr oder Landung, wenn eine Grenze erreicht ist (z. B. Restenergie).'],
]
table('Glossar', 'Glossar', 'Begriffe der App und des Plans.', ['Begriff', 'Erklärung'], [28, 110], G)

# =============================== Änderungen =====================================================
C = [['0.1', '13.10.2026', 'Erste Fassung aus den Vereinbarungen im Chat „Luftschiff_Simulator“: Stichpunkte des Nutzers als Fahrplan der Entscheidungen, feste Auswahl (9 Modelle), Gertler-Formen mit spitzem Heck, Leitwerk vorerst nach Burgess, Vorschläge mit freier Abweichung, Eckdaten immer sichtbar, Speicher-Code, 3D-Vorschau, RC-Blimps und autonome Schiffe, Fahrt mit Assistenten und Missionsplanung, Wetter ruhig/mittel/extrem und historisches Datum, Zeitraffer bis ×1000, Seitenansicht unter 100 m.', 'Claude']]
table('Änderungen', 'Änderungen', 'Versionen dieses Plans.', ['Version', 'Datum', 'Änderung', 'Von'], [9, 12, 120, 10], C)

wb.save(OUT)
print('gespeichert:', OUT)
